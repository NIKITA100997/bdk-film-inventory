import io
from datetime import date

import pytest
from openpyxl import Workbook

from app.services.import_template import _NAME_SERIES_RE
from app.services.invoice_1c import invoice_date_of, invoice_no_from_filename, read_invoice, read_invoice_date


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_invoice_no_from_filename():
    assert invoice_no_from_filename("309-СКУ(М).xls") == "309-СКУ(М)"
    assert invoice_no_from_filename("309-СКУ(М) (1).xls") == "309-СКУ(М)"
    assert invoice_no_from_filename("C:\\fakepath\\412.XLSX") == "412"


def test_read_invoice_rows():
    data = _xlsx([
        ["Счёт на оплату"],
        ["Номенклатура", "Количество"],
        ["С-2 400х2000 - ПЭТ Бежевый (cream silk) (стекло Сатин)", 2.0],
        [None, None],
        ["В-15  600х2000 - Полипропилен Аляска", "30"],
        ["Итого", 32],
    ])
    lines, problems = read_invoice(data, "309.xlsx")
    assert lines == [("С-2 400х2000 - ПЭТ Бежевый (cream silk) (стекло Сатин)", 2), ("В-15 600х2000 - Полипропилен Аляска", 30)]
    assert problems == []


def test_read_invoice_bad_qty_and_no_header():
    lines, problems = read_invoice(_xlsx([["Номенклатура", "Количество"], ["А-1 600х2000", 1.5]]), "1.xlsx")
    assert lines == [] and "не целое" in problems[0]
    with pytest.raises(ValueError, match="Номенклатура"):
        read_invoice(_xlsx([["что-то", "ещё"]]), "1.xlsx")
    with pytest.raises(ValueError, match="прочитать"):
        read_invoice(b"not excel", "1.xls")


def test_series_outside_dictionary_is_extracted():
    assert _NAME_SERIES_RE.match("С-2 400х2000 - ПЭТ").group(1) == "С-2"
    assert _NAME_SERIES_RE.match("В-10.2 800х2000").group(1) == "В-10.2"
    assert _NAME_SERIES_RE.match("Н-1 ВО 800х2000").group(1) == "Н-1 ВО"
    assert _NAME_SERIES_RE.match("ПЭТ Бежевый 400х2000") is None


def test_invoice_key_one_invoice_one_order():
    from app.services.invoice_1c import invoice_key

    same = ["1029-ЛД", "1029-лд(М)", "№ 1029-ЛД", " 1029 - ЛД ", "1029-ЛД (М) (1)"]
    assert {invoice_key(s) for s in same} == {"1029ЛД"}
    assert invoice_key("1029-LД") != invoice_key("1029-ЛД")  # латинская L на Л не похожа — остаётся; A/B/C… — кириллицей
    assert invoice_key("309-CKУ(M)") == invoice_key("309-СКУ")
    assert invoice_key("1029") != invoice_key("1029-ЛД")
    assert invoice_key("1125ЛД") == invoice_key("1125-ЛД") == invoice_key("1125 - лд")


@pytest.mark.parametrize(
    "text,expected",
    [
        ("Счет на оплату № 309 от 16 октября 2026 г.", date(2026, 10, 16)),
        ("Счёт на оплату № 2080-ТДС(М) от 5 Сентября 2026 г.", date(2026, 9, 5)),
        ("Счет № 1 от 07.10.26", date(2026, 10, 7)),
        ("от 31 февраля 2026", None),
        ("Номенклатура Количество", None),
    ],
)
def test_invoice_date_of(text, expected):
    assert invoice_date_of(text) == expected


def test_read_invoice_date_from_header():
    wb = Workbook()
    ws = wb.active
    ws.append(["Счет на оплату № 309 от 16 октября 2026 г."])
    ws.append(["Номенклатура", "Количество"])
    ws.append(["А-1 800х2000 - Манхэттен", 2])
    buf = io.BytesIO()
    wb.save(buf)
    assert read_invoice_date(buf.getvalue(), "309.xlsx") == date(2026, 10, 16)
    # выгрузка без шапки — даты нет, не ошибка
    wb2 = Workbook()
    wb2.active.append(["Номенклатура", "Количество"])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    assert read_invoice_date(buf2.getvalue(), "309.xlsx") is None
