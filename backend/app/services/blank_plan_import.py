"""Раздел про импорт плана заготовок (Excel) — печатная форма («лист
плана окутки/раскроя на дату») даёт то же самое, что наряд-заказ
(services/naryad_import.py) — форму деталей + количество, но, в отличие
от него, ещё и ЦВЕТ построчно (у наряд-заказа плёнка одна на всё
задание — оператор выбирает её отдельно на фронтенде). Поэтому подбор
позиции материала здесь тоже построчный, и делается уже здесь (не на
фронтенде): для каждой строки, если у распознанного цвета ровно одна
активная позиция номенклатуры — используем её материал/толщину как
подсказку, иначе оставляем пустым для ручного выбора на фронтенде.

На одном листе бывает 2-3 таких таблицы бок о бок (разные заказы/даты
печати) — каждая имеет свою колонку "Номенклатура"/"Цвет"/"Заказ", но
раскладка колонок между ними (и между листами/файлами) не совпадает —
как и в naryad_import.py, ищем колонки по тексту заголовка, а не по
номеру. "Номенклатура" объединена (merge) на несколько строк подряд,
когда у одной заготовки несколько цветов — раскрывается в
_resolve_merged до чтения строк, дальше слой парсинга работает с уже
"плоской" сеткой, как и парсер наряд-заказа.

Два слоя, как в naryad_import.py: parse_blank_plan_block/find_blocks —
чистые функции над готовой сеткой (тестируются без файла),
parse_blank_plan_xlsx_bytes — тонкая обёртка чтения .xlsx (openpyxl, не
xlrd — тот .xlsx не читает). enrich_blank_plan_blocks — отдельно,
единственное место, где нужна БД (подбор Part/Color/MaterialSku)."""

import difflib
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session, joinedload

from app.models.dictionaries import Color, MaterialSku, Part, Thickness

HEADER_NOMENKLATURA = "номенклатура"
HEADER_CVET = "цвет"
HEADER_ZAKAZ = "заказ"
TITLE_STOP_WORDS = ("окутка",)

# Цвет в скобках — не обязательно последнее в строке (реальный образец
# "Стоевая (МежКомн) 36х108х2035 (Bolton Oak) ПАЗ-11" — после цвета в
# скобках ещё идёт паз-профиль снаружи скобок), но это ПОСЛЕДНЯЯ
# скобочная группа в названии — "(МежКомн)" всегда идёт раньше неё.
PAREN_GROUP_RE = re.compile(r"\(([^()]+)\)")


def normalize_part_name(text: str) -> str:
    """Убирает шум, которого нет в справочнике деталей ("Заготовка "
    в начале, "(МежКомн)" где угодно), схлопывает пробелы — само
    совпадение со справочником всё равно нечёткое (см.
    enrich_blank_plan_blocks), это только сокращает дистанцию."""
    t = re.sub(r"(?i)^\s*заготовка\s+", "", text)
    t = re.sub(r"(?i)\(\s*межкомн\s*\)", "", t)
    return re.sub(r"\s+", " ", t).strip()


@dataclass(frozen=True)
class ParsedBlankPlanLine:
    part_name_raw: str
    color_raw: str
    quantity_pieces: float


@dataclass(frozen=True)
class BlankPlanBlock:
    sheet_name: str
    suggested_name: str
    lines: list[ParsedBlankPlanLine] = field(default_factory=list)


@dataclass(frozen=True)
class EnrichedBlankPlanLine:
    part_name: str
    suggested_part_id: int | None
    width_mm: float | None
    length_m: float | None
    strip_width_mm: float | None
    color_raw: str
    suggested_sku_id: int | None
    material: str | None
    thickness: float | None
    quantity_pieces: float
    # Раздел про обратную связь — когда однозначно подобрать не удалось
    # (см. enrich_blank_plan_blocks), несколько похожих по сочетанию
    # материал+цвет позиций номенклатуры, чтобы оператор выбрал из них на
    # фронтенде вместо того, чтобы искать заново руками с нуля.
    sku_candidates: list[dict] = field(default_factory=list)


@dataclass(frozen=True)
class EnrichedBlankPlanBlock:
    sheet_name: str
    suggested_name: str
    lines: list[EnrichedBlankPlanLine]


def _cell(row: list, col: int):
    return row[col] if 0 <= col < len(row) else None


def _find_col(row: list, col_start: int, col_end: int, header_text: str) -> int | None:
    for c in range(col_start, min(col_end + 1, len(row))):
        cell = row[c]
        if isinstance(cell, str) and cell.strip().lower() == header_text:
            return c
    return None


def find_blocks(grid: list[list]) -> list[tuple[int, int, int]]:
    """(header_row_idx, col_start, col_end) для каждого блока —
    "Номенклатура" повторяется много раз по вертикали (разрыв печатной
    страницы) на ОДНОЙ и той же стартовой колонке блока — берём только
    первое (самое верхнее) вхождение на каждую колонку, остальные —
    просто повторные заголовки внутри того же блока (см.
    parse_blank_plan_block). Конец блока — колонка перед стартом
    следующего, либо конец сетки для последнего."""
    max_cols = max((len(row) for row in grid), default=0)
    first_header_row_by_col: dict[int, int] = {}
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if isinstance(cell, str) and cell.strip().lower() == HEADER_NOMENKLATURA and c not in first_header_row_by_col:
                first_header_row_by_col[c] = r

    starts = sorted(first_header_row_by_col)
    blocks: list[tuple[int, int, int]] = []
    for i, col_start in enumerate(starts):
        col_end = starts[i + 1] - 1 if i + 1 < len(starts) else max_cols - 1
        blocks.append((first_header_row_by_col[col_start], col_start, col_end))
    return blocks


def _block_title(grid: list[list], header_row_idx: int, col_start: int, col_end: int, sheet_name: str, block_index: int) -> str:
    """Название блока — из строки НАД заголовком колонок (там в реальных
    файлах сидит либо дата запуска, либо текст вида "Окутка МК заказ
    24.08"); если ничего не нашлось — просто лист+номер блока."""
    if header_row_idx - 1 >= 0:
        row_above = grid[header_row_idx - 1]
        for c in range(col_start, min(col_end + 1, len(row_above))):
            cell = row_above[c]
            if isinstance(cell, datetime):
                return f"Заготовки на {cell.date().isoformat()}"
            if isinstance(cell, date):
                return f"Заготовки на {cell.isoformat()}"
            if isinstance(cell, str) and cell.strip():
                return cell.strip()
    return f"{sheet_name} — блок {block_index + 1}"


def parse_blank_plan_block(
    grid: list[list], header_row_idx: int, col_start: int, col_end: int, sheet_name: str, block_index: int
) -> BlankPlanBlock:
    suggested_name = _block_title(grid, header_row_idx, col_start, col_end, sheet_name, block_index)

    header_row = grid[header_row_idx] if header_row_idx < len(grid) else []
    color_col = _find_col(header_row, col_start, col_end, HEADER_CVET)
    order_col = _find_col(header_row, col_start, col_end, HEADER_ZAKAZ)
    if order_col is None:
        # Нестандартная раскладка этого конкретного блока — колонки не
        # нашли, разобрать нечего (не ошибка — остальные блоки листа
        # разбираются независимо).
        return BlankPlanBlock(sheet_name=sheet_name, suggested_name=suggested_name, lines=[])

    lines: list[ParsedBlankPlanLine] = []
    for r in range(header_row_idx + 1, len(grid)):
        row = grid[r]
        nomenklatura_cell = _cell(row, col_start)
        if isinstance(nomenklatura_cell, str) and nomenklatura_cell.strip().lower() == HEADER_NOMENKLATURA:
            continue  # повторный заголовок — разрыв страницы, не новый блок

        qty = _cell(row, order_col)
        if not isinstance(qty, (int, float)) or qty <= 0:
            continue

        name_text = nomenklatura_cell.strip() if isinstance(nomenklatura_cell, str) else ""
        if not name_text or any(w in name_text.lower() for w in TITLE_STOP_WORDS):
            continue

        if color_col is not None:
            color_cell = _cell(row, color_col)
            color_text = color_cell.strip() if isinstance(color_cell, str) else ""
            part_name_raw = name_text
        else:
            matches = list(PAREN_GROUP_RE.finditer(name_text))
            if not matches:
                continue
            last = matches[-1]
            color_text = last.group(1).strip()
            part_name_raw = re.sub(r"\s+", " ", name_text[: last.start()] + name_text[last.end() :]).strip()

        if not color_text:
            continue

        lines.append(
            ParsedBlankPlanLine(
                part_name_raw=normalize_part_name(part_name_raw), color_raw=color_text, quantity_pieces=float(qty)
            )
        )

    return BlankPlanBlock(sheet_name=sheet_name, suggested_name=suggested_name, lines=lines)


def _resolve_merged(ws: Worksheet) -> list[list]:
    """Обычная (не разреженная) сетка значений листа с уже "развёрнутыми"
    объединёнными диапазонами — значение верхней левой ячейки копируется
    во все ячейки диапазона (иначе openpyxl отдаёт None для всех ячеек
    объединения, кроме первой — а "Номенклатура" объединена именно так,
    на несколько строк подряд, когда у детали несколько цветов)."""
    grid = [[cell.value for cell in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]
    for merged_range in ws.merged_cells.ranges:
        top_value = grid[merged_range.min_row - 1][merged_range.min_col - 1]
        for r in range(merged_range.min_row - 1, merged_range.max_row):
            for c in range(merged_range.min_col - 1, merged_range.max_col):
                grid[r][c] = top_value
    return grid


def parse_blank_plan_xlsx_bytes(data: bytes) -> list[BlankPlanBlock]:
    """По всем листам книги сразу (файл — многолетний журнал, лист на
    дату, их бывают десятки) — выбор нужного(ых) листа/блока делает
    оператор уже в предпросмотре на фронтенде."""
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=True)
    blocks: list[BlankPlanBlock] = []
    for ws in workbook.worksheets:
        grid = _resolve_merged(ws)
        for i, (header_row_idx, col_start, col_end) in enumerate(find_blocks(grid)):
            block = parse_blank_plan_block(grid, header_row_idx, col_start, col_end, ws.title, i)
            if block.lines:
                blocks.append(block)
    return blocks


_SKU_MATCH_CUTOFF = 0.45
_SKU_CANDIDATES_MAX = 5


def _sku_label(sku: MaterialSku) -> str:
    return f"{sku.material.name}, {sku.color.name}, {float(sku.thickness.value_mm)} мм"


def enrich_blank_plan_blocks(db: Session, blocks: list[BlankPlanBlock]) -> list[EnrichedBlankPlanBlock]:
    """Подбор детали/материала — единственное место во всём модуле, где
    нужна БД. Деталь — нечёткое совпадение (difflib, стандартная
    библиотека) по нормализованному имени.

    Материал — сначала точное совпадение цвета (как раньше, для файлов,
    где в колонке "Цвет" — только название цвета), и только если под него
    ровно одна активная позиция номенклатуры. Если так не получилось (то
    и другое встречается на практике — раздел обратной связи: в тексте
    бывает написано "ПЭТ Белый", то есть материал+цвет вместе, а не
    отдельно цвет) — нечёткий поиск по сочетанию материал+цвет среди ВСЕХ
    активных позиций: одно уверенное совпадение проставляется тем же
    способом, что и раньше; несколько похожих — не гадаем, какое из них,
    а отдаём списком (sku_candidates) во фронтенд, чтобы оператор выбрал
    сам одним кликом вместо поиска с нуля."""
    parts = db.query(Part).filter(Part.is_active).all()
    part_by_normalized: dict[str, Part] = {p.name.strip().lower(): p for p in parts}
    normalized_names = list(part_by_normalized)

    colors = db.query(Color).all()
    color_by_normalized: dict[str, Color] = {re.sub(r"\s+", " ", c.name.strip().lower()): c for c in colors}

    # thickness > 0 — отсекает позиции-заглушки (материал "Неизвестно",
    # толщина 0) от подсказки: без этого условия такая заглушка (если она
    # единственная активная позиция под цвет) молча подставлялась бы как
    # "уверенное" совпадение — 0 мм всё равно не настоящая толщина, и
    # ProductionTaskLineManualCreate.thickness её никак не отсекает
    # (в отличие от width_mm/length_m там нет Field(gt=0)).
    skus = (
        db.query(MaterialSku)
        .join(Thickness, MaterialSku.thickness_id == Thickness.id)
        .options(joinedload(MaterialSku.material), joinedload(MaterialSku.color), joinedload(MaterialSku.thickness))
        .filter(MaterialSku.is_active, Thickness.value_mm > 0)
        .all()
    )
    skus_by_color_id: dict[int, list[MaterialSku]] = {}
    combined_label_to_sku: dict[str, MaterialSku] = {}
    for s in skus:
        skus_by_color_id.setdefault(s.color_id, []).append(s)
        combined = re.sub(r"\s+", " ", f"{s.material.name} {s.color.name}".strip().lower())
        combined_label_to_sku[combined] = s
    combined_labels = list(combined_label_to_sku)

    result: list[EnrichedBlankPlanBlock] = []
    for block in blocks:
        enriched_lines: list[EnrichedBlankPlanLine] = []
        for line in block.lines:
            normalized = line.part_name_raw.strip().lower()
            match = difflib.get_close_matches(normalized, normalized_names, n=1, cutoff=0.5)
            part = part_by_normalized[match[0]] if match else None

            color_key = re.sub(r"\s+", " ", line.color_raw.strip().lower())
            color = color_by_normalized.get(color_key)
            exact_candidates = skus_by_color_id.get(color.id) if color else None
            sku = exact_candidates[0] if exact_candidates and len(exact_candidates) == 1 else None

            sku_candidates: list[dict] = []
            if sku is None:
                fuzzy_labels = difflib.get_close_matches(
                    color_key, combined_labels, n=_SKU_CANDIDATES_MAX, cutoff=_SKU_MATCH_CUTOFF
                )
                fuzzy_skus = [combined_label_to_sku[label] for label in fuzzy_labels]
                if len(fuzzy_skus) == 1:
                    sku = fuzzy_skus[0]
                elif len(fuzzy_skus) > 1:
                    # get_close_matches уже отдаёт по убыванию похожести —
                    # если лучший результат заметно впереди второго (не
                    # просто "тоже похоже, но не факт какой из двух"),
                    # берём его не гадая: комбинация материал+цвет обычно
                    # даёт именно такой явный отрыв, в отличие от случая
                    # "два реальных близких варианта" (напр. одно и то же
                    # название цвета у ПЭТ 2Д и ПЭТ 3Д) — там разница
                    # ratio() между лучшим и вторым мала, это и остаётся
                    # списком на выбор.
                    ratios = [difflib.SequenceMatcher(None, color_key, label).ratio() for label in fuzzy_labels]
                    if ratios[0] >= 0.92 or ratios[0] - ratios[1] >= 0.08:
                        sku = fuzzy_skus[0]
                    else:
                        sku_candidates = [{"sku_id": s.id, "label": _sku_label(s)} for s in fuzzy_skus]

            enriched_lines.append(
                EnrichedBlankPlanLine(
                    part_name=part.name if part else line.part_name_raw,
                    suggested_part_id=part.id if part else None,
                    width_mm=float(part.width_mm) if part else None,
                    length_m=float(part.length_m) if part else None,
                    strip_width_mm=float(part.strip_width_mm) if part and part.strip_width_mm is not None else None,
                    color_raw=line.color_raw,
                    suggested_sku_id=sku.id if sku else None,
                    material=sku.material.name if sku else None,
                    thickness=float(sku.thickness.value_mm) if sku else None,
                    quantity_pieces=line.quantity_pieces,
                    sku_candidates=sku_candidates,
                )
            )
        result.append(
            EnrichedBlankPlanBlock(sheet_name=block.sheet_name, suggested_name=block.suggested_name, lines=enriched_lines)
        )
    return result
