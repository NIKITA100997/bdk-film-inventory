"""Выгрузка счёта из 1С (16.10): .xls/.xlsx с колонками «Номенклатура» и
«Количество», имя файла — номер счёта («309-СКУ(М).xls»). Превращаем в
строки «наименование<TAB>кол-во» — тот же текст, что вставляют вручную в
окно «Из счёта 1С или графика», дальше обычный разбор по шаблону типа."""

import io
from datetime import date
import re

import xlrd

_NAME_HEADERS = ("номенклатура", "наименование", "товар", "товары (работы, услуги)")
_QTY_HEADERS = ("количество", "кол-во", "кол.")
_COPY_SUFFIX_RE = re.compile(r"\s+\(\d+\)$")  # «309-СКУ(М) (1).xls» — вторая загрузка браузером


_LAT2CYR = str.maketrans("ABCEHKMOPTXY", "АВСЕНКМОРТХУ")


def invoice_key(no: str | None) -> str:
    """Один счёт — один заказ клиента (20.10): «1029-ЛД», «1029-лд(М)»,
    «№ 1029-ЛД» — один ключ «1029ЛД». Без «№», пробелов, дефисов, пометок в скобках
    в конце, регистра; латиница, похожая на кириллицу, — кириллицей."""
    s = (no or "").upper().replace("№", "")
    s = re.sub(r"(\s*\([^)]*\))+\s*$", "", s)  # «(М)» и подобные пометки в конце
    # без пробелов и дефисов: «1125ЛД» = «1125-ЛД» = «1836 - АТ»
    return re.sub(r"[\s\-–—]+", "", s).translate(_LAT2CYR)


def invoice_no_from_filename(filename: str) -> str:
    stem = re.sub(r"\.(xlsx?|xls)$", "", (filename or "").strip().replace("\\", "/").split("/")[-1], flags=re.I)
    return _COPY_SUFFIX_RE.sub("", stem).strip()


def _rows(data: bytes, filename: str) -> list[list[object]]:
    if (filename or "").lower().endswith(".xlsx") or data[:2] == b"PK":
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    sh = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
    return [sh.row_values(i) for i in range(sh.nrows)]


def _qty(v: object) -> int | None:
    if v is None or v == "":
        return None
    try:
        q = float(str(v).replace(",", ".").replace("\xa0", "").replace(" ", ""))
    except ValueError:
        return None
    return int(q) if q > 0 and q == int(q) else None


def read_invoice(data: bytes, filename: str) -> tuple[list[tuple[str, int]], list[str]]:
    """Строки счёта (наименование, кол-во) и замечания. Шапку ищем по
    названиям колонок; «Итого» и пустые строки пропускаем."""
    try:
        rows = _rows(data, filename)
    except Exception as e:  # noqa: BLE001 — любой битый файл → понятная ошибка
        raise ValueError("Не удалось прочитать файл — нужна выгрузка счёта из 1С в Excel (.xls или .xlsx)") from e
    i_name = i_qty = None
    start = 0
    for n, r in enumerate(rows):
        cells = [str(c or "").strip().lower() for c in r]
        names = [i for i, c in enumerate(cells) if c in _NAME_HEADERS]
        qtys = [i for i, c in enumerate(cells) if c in _QTY_HEADERS]
        if names and qtys:
            i_name, i_qty, start = names[0], qtys[0], n + 1
            break
    if i_name is None:
        raise ValueError("В файле нет колонок «Номенклатура» и «Количество»")
    out: list[tuple[str, int]] = []
    problems: list[str] = []
    for n, r in enumerate(rows[start:], start=start + 1):
        name = " ".join(str(r[i_name] if i_name < len(r) and r[i_name] is not None else "").split())
        if not name or name.lower().startswith("итого"):
            continue
        q = _qty(r[i_qty] if i_qty < len(r) else None)
        if q is None:
            problems.append(f"Строка {n}: «{name}» — количество не целое число")
            continue
        out.append((name, q))
    if not out and not problems:
        problems.append("В счёте нет строк с позициями")
    return out, problems


_MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}
_DATE_WORDS_RE = re.compile(r"\bот\s+(\d{1,2})\s+([а-яё]+)\s+(\d{4})", re.I)
_DATE_DOTS_RE = re.compile(r"\bот\s+(\d{1,2})\.(\d{1,2})\.(\d{2,4})\b", re.I)


def invoice_date_of(text: str) -> date | None:
    """Дата из «Счет на оплату № 309 от 16 октября 2026 г.» или «… от 16.10.2026»."""
    for m in _DATE_WORDS_RE.finditer(text or ""):
        month = _MONTHS.get(m.group(2).lower().replace("ё", "е"))
        if month:
            try:
                return date(int(m.group(3)), month, int(m.group(1)))
            except ValueError:
                continue
    for m in _DATE_DOTS_RE.finditer(text or ""):
        y = int(m.group(3))
        try:
            return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
        except ValueError:
            continue
    return None


def read_invoice_date(data: bytes, filename: str) -> date | None:
    """Дата счёта из шапки выгрузки 1С (28.10) — первые строки до таблицы."""
    try:
        rows = _rows(data, filename)
    except Exception:  # noqa: BLE001 — нет даты, не ошибка: строки читает read_invoice
        return None
    for r in rows[:30]:
        d = invoice_date_of(" ".join(str(c) for c in r if c not in (None, "")))
        if d:
            return d
    return None
